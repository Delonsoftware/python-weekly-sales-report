#!/usr/bin/env python3
"""
Automated Weekly Sales Report Generator
=======================================

Turns a raw sales CSV into a ready-to-send weekly report in one command:

    output/weekly_sales_report_<week>.xlsx   formatted workbook with charts
    output/weekly_summary_<week>.txt         plain-English summary (paste into an email)

The report covers one Monday-to-Sunday week and compares it with the week before:
revenue, orders, units, average order value, refunds, daily pattern, top products,
categories, regions, sales reps, an 8-week trend, and the biggest movers.

Rows with missing or invalid data are skipped and listed on their own sheet, so
nothing is silently ignored.

Usage
-----
    python generate_sample_sales.py                   # make sample data first
    python sales_report.py                            # latest complete week
    python sales_report.py --week-start 2026-09-28    # a specific week (any day in it works)
    python sales_report.py --input my_sales.csv --output-dir reports

Required CSV columns: order_id, order_date, product, quantity, unit_price
Optional columns:     category, region, sales_rep, discount_pct, status
(status values: Completed, Refunded, Cancelled; blank counts as Completed)
"""

import argparse
import logging
import math
import sys
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

log = logging.getLogger("sales_report")

REQUIRED_COLS = ["order_id", "order_date", "product", "quantity", "unit_price"]
OPTIONAL_COLS = ["category", "region", "sales_rep", "discount_pct", "status"]
MIN_PRIOR_REVENUE = 250.0      # ignore tiny products when naming "biggest movers"
NOTABLE_MOVE = 0.10            # a mover must be at least +/-10%
REFUND_ALERT_RATE = 0.06


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------
def money(value: float, decimals: int = 0) -> str:
    return f"${value:,.{decimals}f}"


def pct_change(current, previous):
    """Fractional change, or None when there is nothing to compare with."""
    if previous is None or pd.isna(previous) or previous == 0:
        return None
    return (current - previous) / previous


def pct_text(change) -> str:
    return "n/a" if change is None else f"{change * 100:+.1f}%"


# --------------------------------------------------------------------------
# Loading and validating
# --------------------------------------------------------------------------
def load_sales(path: Path):
    """Read the CSV; return (valid rows, skipped rows with reasons)."""
    if not path.exists():
        sys.exit(f"Input file not found: {path}\n"
                 "Run 'python generate_sample_sales.py' first to create sample data.")

    raw = pd.read_csv(path, dtype=str, keep_default_na=False)
    raw.columns = [c.strip().lower().replace(" ", "_") for c in raw.columns]
    missing = [c for c in REQUIRED_COLS if c not in raw.columns]
    if missing:
        sys.exit(f"Input file is missing required columns: {', '.join(missing)}")
    for col in OPTIONAL_COLS:
        if col not in raw.columns:
            raw[col] = ""

    raw.insert(0, "source_row", range(2, len(raw) + 2))
    df = raw.copy()
    for col in ["order_id", "product", "category", "region", "sales_rep", "status"]:
        df[col] = df[col].str.strip()

    df["order_date"] = pd.to_datetime(df["order_date"], errors="coerce", format="mixed")
    df["quantity"] = pd.to_numeric(df["quantity"].str.replace(",", "", regex=False),
                                   errors="coerce")
    df["unit_price"] = pd.to_numeric(
        df["unit_price"].str.replace(r"[$,\s]", "", regex=True), errors="coerce")
    discount = df["discount_pct"].str.replace("%", "", regex=False).str.strip()
    df["discount_pct"] = pd.to_numeric(discount.replace("", "0"), errors="coerce")

    checks = [
        ("Missing order ID", df["order_id"] == ""),
        ("Missing product", df["product"] == ""),
        ("Invalid date", df["order_date"].isna()),
        ("Invalid quantity", df["quantity"].isna() | (df["quantity"] <= 0)),
        ("Invalid price", df["unit_price"].isna() | (df["unit_price"] < 0)),
        ("Invalid discount", df["discount_pct"].isna()
         | (df["discount_pct"] < 0) | (df["discount_pct"] > 100)),
    ]
    reason = pd.Series("", index=df.index)
    for label, mask in checks:
        reason[(reason == "") & mask] = label

    skipped = raw[reason != ""].copy()
    skipped.insert(1, "reason", reason[reason != ""])

    valid = df[reason == ""].copy()
    for col in ["category", "region", "sales_rep"]:
        valid[col] = valid[col].replace("", "Unknown")

    status = valid["status"].str.lower()
    valid["status"] = "completed"
    valid.loc[status.isin(["refunded", "refund", "returned"]), "status"] = "refunded"
    valid.loc[status.isin(["cancelled", "canceled"]), "status"] = "cancelled"

    valid["revenue"] = (valid["quantity"] * valid["unit_price"]
                        * (1 - valid["discount_pct"] / 100)).round(2)
    valid["order_day"] = valid["order_date"].dt.normalize()
    return valid.reset_index(drop=True), skipped.reset_index(drop=True)


# --------------------------------------------------------------------------
# Calculations
# --------------------------------------------------------------------------
def pick_week_start(df: pd.DataFrame, requested: str = None) -> pd.Timestamp:
    """Monday of the requested week, or of the latest *complete* week in the data."""
    if requested:
        day = pd.Timestamp(requested).normalize()
        return day - timedelta(days=day.weekday())
    last = df["order_day"].max()
    start = last - timedelta(days=last.weekday())
    if last.weekday() != 6:          # data stops mid-week: use the week before
        start -= timedelta(days=7)
    return start


def slice_week(df: pd.DataFrame, start: pd.Timestamp) -> pd.DataFrame:
    return df[(df["order_day"] >= start) & (df["order_day"] < start + timedelta(days=7))]


def period_stats(df: pd.DataFrame) -> dict:
    done = df[df["status"] == "completed"]
    refunded = df[df["status"] == "refunded"]
    orders = int(done["order_id"].nunique())
    refund_orders = int(refunded["order_id"].nunique())
    revenue = round(float(done["revenue"].sum()), 2)
    return {
        "revenue": revenue,
        "orders": orders,
        "units": int(done["quantity"].sum()),
        "aov": revenue / orders if orders else 0.0,
        "refund_orders": refund_orders,
        "refund_amount": round(float(refunded["revenue"].sum()), 2),
        "refund_rate": refund_orders / (orders + refund_orders) if (orders + refund_orders) else 0.0,
    }


def daily_table(week: pd.DataFrame, prev: pd.DataFrame, start: pd.Timestamp) -> pd.DataFrame:
    days = pd.date_range(start, periods=7)
    done, prev_done = week[week["status"] == "completed"], prev[prev["status"] == "completed"]
    revenue = done.groupby("order_day")["revenue"].sum().reindex(days, fill_value=0.0)
    orders = done.groupby("order_day")["order_id"].nunique().reindex(days, fill_value=0)
    units = done.groupby("order_day")["quantity"].sum().reindex(days, fill_value=0)
    last_week = (prev_done.groupby("order_day")["revenue"].sum()
                 .reindex(days - timedelta(days=7), fill_value=0.0))
    return pd.DataFrame({
        "Date": [d.date() for d in days],
        "Day": days.strftime("%a"),
        "Revenue": revenue.round(2).values,
        "Orders": orders.astype(int).values,
        "Units": units.astype(int).values,
        "Last Week Revenue": last_week.round(2).values,
    })


def breakdown(cur_done: pd.DataFrame, prev_done: pd.DataFrame, col: str, label: str) -> pd.DataFrame:
    grouped = cur_done.groupby(col).agg(
        revenue=("revenue", "sum"), units=("quantity", "sum"),
        orders=("order_id", "nunique"))
    prev = prev_done.groupby(col)["revenue"].sum().rename("prev_revenue")
    out = grouped.join(prev, how="left")
    out["prev_revenue"] = out["prev_revenue"].fillna(0.0)
    total = out["revenue"].sum()
    out["share"] = out["revenue"] / total if total else 0.0
    out["change_pct"] = [pct_change(r, p) for r, p in zip(out["revenue"], out["prev_revenue"])]
    out["revenue"] = out["revenue"].round(2)
    out["prev_revenue"] = out["prev_revenue"].round(2)
    out = out.sort_values("revenue", ascending=False).reset_index()
    out = out.rename(columns={col: label})
    return out[[label, "revenue", "share", "units", "orders", "prev_revenue", "change_pct"]]


def trend_table(df: pd.DataFrame, start: pd.Timestamp, weeks: int = 8) -> pd.DataFrame:
    first_day = df["order_day"].min()
    first_monday = first_day - timedelta(days=first_day.weekday())
    starts = [start - timedelta(days=7 * i) for i in range(weeks - 1, -1, -1)]
    rows = []
    for week_start in (s for s in starts if s >= first_monday):
        done = slice_week(df, week_start)
        done = done[done["status"] == "completed"]
        rows.append({
            "Week Starting": week_start.strftime("%b %d"),
            "Revenue": round(float(done["revenue"].sum()), 2),
            "Orders": int(done["order_id"].nunique()),
        })
    return pd.DataFrame(rows)


def find_movers(products: pd.DataFrame, min_prior: float):
    eligible = products[(products["prev_revenue"] >= min_prior) & products["change_pct"].notna()]
    if eligible.empty:
        return None, None
    best = eligible.loc[eligible["change_pct"].astype(float).idxmax()]
    worst = eligible.loc[eligible["change_pct"].astype(float).idxmin()]
    return (best if best["change_pct"] >= NOTABLE_MOVE else None,
            worst if worst["change_pct"] <= -NOTABLE_MOVE else None)


# --------------------------------------------------------------------------
# Plain-English summary
# --------------------------------------------------------------------------
def build_highlights(start, cur, prev, has_prev, daily, products, regions, reps, mover_up, mover_down):
    lines = []
    change = pct_change(cur["revenue"], prev["revenue"]) if has_prev else None

    if change is None:
        lines.append(f"Revenue for the week was {money(cur['revenue'])} from {cur['orders']:,} orders "
                     f"(no prior-week data to compare with).")
    else:
        direction = "up" if change >= 0 else "down"
        lines.append(f"Revenue for the week was {money(cur['revenue'])} from {cur['orders']:,} orders, "
                     f"{direction} {abs(change) * 100:.1f}% from last week ({money(prev['revenue'])}).")

    aov_change = pct_change(cur["aov"], prev["aov"]) if has_prev else None
    lines.append(f"Average order value was {money(cur['aov'], 2)} ({pct_text(aov_change)} vs last week) "
                 f"and {cur['units']:,} units were sold.")

    if cur["revenue"] > 0:
        best_day = daily.loc[daily["Revenue"].idxmax()]
        worst_day = daily.loc[daily["Revenue"].idxmin()]
        lines.append(f"{best_day['Date']:%A} was the strongest day at {money(best_day['Revenue'])}; "
                     f"{worst_day['Date']:%A} was the slowest at {money(worst_day['Revenue'])}.")

        top = products.iloc[0]
        lines.append(f"{top['Product']} was the top seller with {money(top['revenue'])} "
                     f"({top['share']:.0%} of revenue).")

    if mover_up is not None:
        lines.append(f"Biggest gain: {mover_up['Product']}, up {mover_up['change_pct'] * 100:.0f}% "
                     f"to {money(mover_up['revenue'])} (from {money(mover_up['prev_revenue'])}).")
    if mover_down is not None:
        lines.append(f"Biggest drop: {mover_down['Product']}, down {abs(mover_down['change_pct']) * 100:.0f}% "
                     f"to {money(mover_down['revenue'])} (from {money(mover_down['prev_revenue'])}).")

    if len(regions):
        top_region = regions.iloc[0]
        lines.append(f"{top_region['Region']} led all regions with {top_region['share']:.0%} of revenue.")
    if len(reps):
        top_rep = reps.iloc[0]
        lines.append(f"Top sales rep: {top_rep['Sales Rep']} with {money(top_rep['revenue'])} in sales.")

    if cur["refund_orders"] or prev["refund_orders"]:
        text = (f"Refunds: {cur['refund_orders']:,} orders ({money(cur['refund_amount'])}), "
                f"{cur['refund_rate']:.1%} of orders")
        rate_jump = cur["refund_rate"] - prev["refund_rate"] if has_prev else 0
        if cur["refund_rate"] >= REFUND_ALERT_RATE or rate_jump >= 0.02:
            text += f", up from {prev['refund_rate']:.1%} last week. Worth a look."
        else:
            text += "."
        lines.append(text)
    return lines


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GOOD_FONT = Font(color="1E7B34", bold=True)
BAD_FONT = Font(color="C00000", bold=True)
MONEY_FMT = '"$"#,##0.00'
PCT_FMT = "0.0%"


def style_header(ws, row=1):
    for cell in ws[row]:
        if cell.value is not None:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_table(ws, df, formats=None, widths=None, change_cols=(), empty_note=None):
    """Write a DataFrame as a styled table starting at A1."""
    formats = formats or {}
    ws.append(list(df.columns))
    style_header(ws)
    for row in df.itertuples(index=False):
        values = []
        for value in row:
            if value is None or (isinstance(value, float) and math.isnan(value)):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            values.append(value)
        ws.append(values)

    for idx, name in enumerate(df.columns, start=1):
        fmt = formats.get(name)
        for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
            cell = cells[0]
            if name in change_cols:
                if cell.value is None:
                    cell.value, cell.alignment = "n/a", Alignment(horizontal="right")
                else:
                    cell.number_format = "+0.0%;-0.0%;0.0%"
                    cell.font = GOOD_FONT if cell.value > 0 else BAD_FONT if cell.value < 0 else Font()
            elif fmt:
                cell.number_format = fmt
    for idx, width in (widths or {}).items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    elif empty_note:
        ws["A2"] = empty_note


def breakdown_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns={
        "revenue": "Revenue", "share": "% of Revenue", "units": "Units",
        "orders": "Orders", "prev_revenue": "Last Week Revenue",
        "change_pct": "Change vs Last Week",
    })


BREAKDOWN_FORMATS = {"Revenue": MONEY_FMT, "% of Revenue": PCT_FMT,
                     "Last Week Revenue": MONEY_FMT, "Units": "#,##0", "Orders": "#,##0"}
BREAKDOWN_WIDTHS = {1: 30, 2: 14, 3: 13, 4: 10, 5: 10, 6: 19, 7: 20}


def build_summary_sheet(ws, start, end, cur, prev, has_prev, highlights, info):
    ws["A1"] = f"Weekly Sales Report: {start:%a %b %d} to {end:%a %b %d, %Y}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"Generated {datetime.now():%Y-%m-%d %H:%M}  |  "
                f"Source: {info['source']}  |  Rows used: {info['rows_used']:,}  |  "
                f"Rows skipped (invalid): {info['rows_skipped']:,}")
    ws["A2"].font = Font(italic=True, color="595959")

    ws.append([])
    ws.append(["Metric", "This Week", "Last Week", "Change", "% Change"])
    style_header(ws, row=4)

    kpis = [  # label, key, number format, is an increase good?
        ("Revenue", "revenue", MONEY_FMT, True),
        ("Orders", "orders", "#,##0", True),
        ("Units sold", "units", "#,##0", True),
        ("Average order value", "aov", MONEY_FMT, True),
        ("Refunded orders", "refund_orders", "#,##0", False),
        ("Refunded amount", "refund_amount", MONEY_FMT, False),
    ]
    for label, key, fmt, up_is_good in kpis:
        row = ws.max_row + 1
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, cur[key]).number_format = fmt
        if has_prev:
            ws.cell(row, 3, prev[key]).number_format = fmt
            ws.cell(row, 4, cur[key] - prev[key]).number_format = fmt
            change = pct_change(cur[key], prev[key])
        else:
            ws.cell(row, 3, "n/a").alignment = Alignment(horizontal="right")
            ws.cell(row, 4, "n/a").alignment = Alignment(horizontal="right")
            change = None
        cell = ws.cell(row, 5, change if change is not None else "n/a")
        if change is None:
            cell.alignment = Alignment(horizontal="right")
        else:
            cell.number_format = "+0.0%;-0.0%;0.0%"
            if change != 0:
                good = (change > 0) == up_is_good
                cell.font = GOOD_FONT if good else BAD_FONT

    row = ws.max_row + 2
    ws.cell(row, 1, "Highlights").font = Font(bold=True, size=13)
    for line in highlights:
        row += 1
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row, 1, "\u2022 " + line)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 16 * max(1, math.ceil(len(line) / 90))

    for idx, width in {1: 26, 2: 16, 3: 16, 4: 14, 5: 12}.items():
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_charts(summary_ws, daily_ws, trend_ws, n_trend_rows):
    bar = BarChart()
    bar.type = "col"
    bar.title = "Daily revenue: this week vs last week"
    bar.add_data(Reference(daily_ws, min_col=3, min_row=1, max_row=8), titles_from_data=True)
    bar.add_data(Reference(daily_ws, min_col=6, min_row=1, max_row=8), titles_from_data=True)
    bar.set_categories(Reference(daily_ws, min_col=2, min_row=2, max_row=8))
    bar.series[0].tx = SeriesLabel(v="This week")
    bar.series[1].tx = SeriesLabel(v="Last week")
    bar.y_axis.title = "Revenue ($)"
    bar.y_axis.numFmt = '"$"#,##0'
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.height, bar.width = 7.5, 16
    summary_ws.add_chart(bar, "G4")

    if n_trend_rows >= 2:
        line = LineChart()
        line.title = "Weekly revenue trend"
        line.add_data(Reference(trend_ws, min_col=2, min_row=1, max_row=n_trend_rows + 1),
                      titles_from_data=True)
        line.set_categories(Reference(trend_ws, min_col=1, min_row=2, max_row=n_trend_rows + 1))
        line.series[0].smooth = False
        line.legend = None
        line.y_axis.title = "Revenue ($)"
        line.y_axis.numFmt = '"$"#,##0'
        line.x_axis.delete = False
        line.y_axis.delete = False
        line.height, line.width = 7.5, 16
        summary_ws.add_chart(line, "G20")


def export_excel(path, start, end, cur, prev, has_prev, highlights, daily, products,
                 categories, regions, reps, trend, skipped, info):
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    daily_ws = wb.create_sheet("Daily")
    products_ws = wb.create_sheet("Top Products")
    categories_ws = wb.create_sheet("Categories")
    regions_ws = wb.create_sheet("Regions")
    reps_ws = wb.create_sheet("Sales Reps")
    trend_ws = wb.create_sheet("Weekly Trend")
    skipped_ws = wb.create_sheet("Skipped Rows")

    write_table(daily_ws, daily,
                formats={"Revenue": MONEY_FMT, "Last Week Revenue": MONEY_FMT,
                         "Date": "yyyy-mm-dd", "Orders": "#,##0", "Units": "#,##0"},
                widths={1: 13, 2: 8, 3: 14, 4: 10, 5: 10, 6: 19})
    for label, ws, table in [("Product", products_ws, products), ("Category", categories_ws, categories),
                             ("Region", regions_ws, regions), ("Sales Rep", reps_ws, reps)]:
        write_table(ws, breakdown_for_excel(table), BREAKDOWN_FORMATS, BREAKDOWN_WIDTHS,
                    change_cols=("Change vs Last Week",))
    write_table(trend_ws, trend, formats={"Revenue": MONEY_FMT, "Orders": "#,##0"},
                widths={1: 16, 2: 14, 3: 10})
    write_table(skipped_ws, skipped, widths={1: 12, 2: 20, 3: 14, 4: 14, 5: 24},
                empty_note="No invalid rows: every row in the file was used.")

    build_summary_sheet(summary_ws, start, end, cur, prev, has_prev, highlights, info)
    summary_ws.page_setup.orientation = "landscape"
    summary_ws.page_setup.fitToWidth = 1
    summary_ws.page_setup.fitToHeight = 1
    summary_ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    add_charts(summary_ws, daily_ws, trend_ws, len(trend))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# --------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------
def run_report(input_path: Path, output_dir: Path, week_start: str = None,
               min_prior: float = MIN_PRIOR_REVENUE) -> dict:
    df, skipped = load_sales(input_path)
    log.info("Loaded %d valid rows (%d skipped as invalid)", len(df), len(skipped))
    if df.empty:
        sys.exit("No valid rows to report on.")

    start = pick_week_start(df, week_start)
    end = start + timedelta(days=6)
    week = slice_week(df, start)
    prev_week = slice_week(df, start - timedelta(days=7))
    if week.empty:
        sys.exit(f"No sales found for the week of {start:%Y-%m-%d}. "
                 f"Data runs from {df['order_day'].min():%Y-%m-%d} to {df['order_day'].max():%Y-%m-%d}.")
    log.info("Reporting on %s to %s", f"{start:%Y-%m-%d}", f"{end:%Y-%m-%d}")

    cur, prev = period_stats(week), period_stats(prev_week)
    has_prev = not prev_week.empty
    done, prev_done = week[week["status"] == "completed"], prev_week[prev_week["status"] == "completed"]

    daily = daily_table(week, prev_week, start)
    products = breakdown(done, prev_done, "product", "Product")
    categories = breakdown(done, prev_done, "category", "Category")
    regions = breakdown(done, prev_done, "region", "Region")
    reps = breakdown(done, prev_done, "sales_rep", "Sales Rep")
    trend = trend_table(df, start)
    mover_up, mover_down = find_movers(products, min_prior)

    highlights = build_highlights(start, cur, prev, has_prev, daily, products,
                                  regions, reps, mover_up, mover_down)
    info = {"source": input_path.name, "rows_used": len(df), "rows_skipped": len(skipped)}

    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = f"{start:%Y-%m-%d}"
    xlsx_path = output_dir / f"weekly_sales_report_{stamp}.xlsx"
    txt_path = output_dir / f"weekly_summary_{stamp}.txt"

    export_excel(xlsx_path, start, end, cur, prev, has_prev, highlights, daily, products,
                 categories, regions, reps, trend, skipped, info)

    heading = f"Weekly Sales Summary: {start:%b %d} to {end:%b %d, %Y}"
    txt_path.write_text(heading + "\n" + "=" * len(heading) + "\n\n"
                        + "\n".join(f"- {line}" for line in highlights) + "\n",
                        encoding="utf-8")

    return {"start": start, "end": end, "current": cur, "previous": prev,
            "highlights": highlights, "daily": daily, "products": products,
            "skipped": skipped, "xlsx": xlsx_path, "txt": txt_path}


def setup_logging() -> None:
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s  %(message)s", "%H:%M:%S")
    console = logging.StreamHandler(sys.stdout)
    console.setFormatter(fmt)
    file_handler = logging.FileHandler("sales_report.log", encoding="utf-8")
    file_handler.setFormatter(fmt)
    log.setLevel(logging.INFO)
    log.addHandler(console)
    log.addHandler(file_handler)


def main() -> int:
    parser = argparse.ArgumentParser(description="Generate a weekly sales report")
    parser.add_argument("--input", type=Path, default=Path("data/sales.csv"))
    parser.add_argument("--output-dir", type=Path, default=Path("output"))
    parser.add_argument("--week-start", default=None,
                        help="any date inside the week to report on (default: latest complete week)")
    parser.add_argument("--min-prior", type=float, default=MIN_PRIOR_REVENUE,
                        help="minimum last-week revenue for a product to count as a 'mover'")
    args = parser.parse_args()

    setup_logging()
    started = datetime.now()
    result = run_report(args.input, args.output_dir, args.week_start, args.min_prior)

    log.info("-" * 60)
    for line in result["highlights"]:
        log.info("* %s", line)
    log.info("-" * 60)
    log.info("Report saved to:  %s", result["xlsx"])
    log.info("Summary saved to: %s", result["txt"])
    log.info("Done in %.1f seconds.", (datetime.now() - started).total_seconds())
    return 0


if __name__ == "__main__":
    sys.exit(main())
