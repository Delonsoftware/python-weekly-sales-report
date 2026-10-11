"""
Tests for the weekly sales report.

Run with either:
    python test_sales_report.py
    pytest -q          (if you have pytest installed)
"""

import tempfile
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

import generate_sample_sales as gen
import sales_report as sr


def make_csv(tmp: Path, rows: list) -> Path:
    path = tmp / "sales.csv"
    pd.DataFrame(rows, columns=gen.FIELDS).to_csv(path, index=False)
    return path


def row(order_id, day, product, qty, price, discount=0, status="Completed",
        category="Gadgets", region="Toronto", rep="Alex"):
    return {"order_id": order_id, "order_date": day, "product": product,
            "category": category, "region": region, "sales_rep": rep,
            "quantity": qty, "unit_price": price, "discount_pct": discount,
            "status": status}


# Mon 2026-09-21 .. Sun 2026-09-27 is "last week"; Mon 2026-09-28 .. Sun 2026-10-04 is "this week"
KNOWN_ROWS = [
    row("A1", "2026-09-22", "Widget", 2, "10.00"),                  # last week: 20.00
    row("A2", "2026-09-23", "Gadget", 1, "100.00"),                 # last week: 100.00
    row("B1", "2026-09-28", "Widget", 3, "10.00"),                  # this week: 30.00
    row("B2", "2026-09-29", "Gadget", 2, "100.00", discount=10),    # this week: 180.00
    row("B3", "2026-09-30", "Widget", 1, "10.00", status="Refunded"),   # refund: not revenue
    row("B4", "2026-10-01", "Gizmo", 1, "50.00", status="Cancelled"),   # ignored entirely
    row("B5", "2026-10-04", "Widget", 1, "10.00"),                  # this week: 10.00
]


def test_week_selection():
    with tempfile.TemporaryDirectory() as tmp:
        df, _ = sr.load_sales(make_csv(Path(tmp), KNOWN_ROWS))
        # data ends on a Sunday -> that week is complete
        assert sr.pick_week_start(df) == pd.Timestamp("2026-09-28")
        # data ends mid-week -> fall back to the previous complete week
        mid = df[df["order_day"] <= "2026-10-01"]
        assert sr.pick_week_start(mid) == pd.Timestamp("2026-09-21")
        # any date inside a week snaps to its Monday
        assert sr.pick_week_start(df, "2026-09-24") == pd.Timestamp("2026-09-21")


def test_revenue_math_and_refunds():
    with tempfile.TemporaryDirectory() as tmp:
        df, _ = sr.load_sales(make_csv(Path(tmp), KNOWN_ROWS))
        this_week = sr.period_stats(sr.slice_week(df, pd.Timestamp("2026-09-28")))
        last_week = sr.period_stats(sr.slice_week(df, pd.Timestamp("2026-09-21")))

        assert this_week["revenue"] == 220.00          # 30 + 180 (10% off) + 10
        assert this_week["orders"] == 3                # B1, B2, B5
        assert this_week["units"] == 6
        assert round(this_week["aov"], 2) == round(220 / 3, 2)
        assert this_week["refund_orders"] == 1         # B3
        assert this_week["refund_amount"] == 10.00
        assert last_week["revenue"] == 120.00
        # cancelled order never shows up anywhere
        assert "Gizmo" not in set(df[df["status"] != "cancelled"]["product"])


def test_percent_change():
    assert sr.pct_change(110, 100) == 0.10
    assert sr.pct_change(50, 100) == -0.50
    assert sr.pct_change(10, 0) is None
    assert sr.pct_change(10, None) is None


def test_invalid_rows_are_skipped_and_reported():
    bad = [
        row("X1", "not a date", "Widget", 1, "10.00"),
        row("X2", "2026-09-28", "Widget", "", "10.00"),
        row("X3", "2026-09-28", "Widget", 1, "-5.00"),
        row("X4", "2026-09-28", "Widget", 1, "10.00", discount=150),
        row("X5", "2026-09-28", "", 1, "10.00"),
    ]
    with tempfile.TemporaryDirectory() as tmp:
        df, skipped = sr.load_sales(make_csv(Path(tmp), KNOWN_ROWS + bad))
        assert len(df) == len(KNOWN_ROWS)
        assert len(skipped) == 5
        assert set(skipped["reason"]) == {"Invalid date", "Invalid quantity", "Invalid price",
                                          "Invalid discount", "Missing product"}


def test_movers_ignore_tiny_products():
    products = pd.DataFrame({
        "Product": ["Big", "Tiny"], "revenue": [900.0, 20.0],
        "prev_revenue": [600.0, 1.0], "change_pct": [0.5, 19.0],
    })
    up, down = sr.find_movers(products, min_prior=250)
    assert up["Product"] == "Big"      # "Tiny" jumped 1900% but from $1, so it's ignored
    assert down is None


def test_end_to_end_report():
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        end = date(2026, 9, 27)                       # a Sunday
        rows = gen.generate(weeks=10, end=end, seed=3)
        csv_path = make_csv(tmp, rows)

        result = sr.run_report(csv_path, tmp / "out")
        assert result["start"] == pd.Timestamp("2026-09-21")
        assert result["xlsx"].exists() and result["txt"].exists()

        # independent check of the headline number straight from the CSV
        raw = pd.read_csv(csv_path)
        raw["order_date"] = pd.to_datetime(raw["order_date"], errors="coerce")
        raw["unit_price"] = pd.to_numeric(raw["unit_price"], errors="coerce")
        raw["quantity"] = pd.to_numeric(raw["quantity"], errors="coerce")
        raw["discount_pct"] = pd.to_numeric(raw["discount_pct"], errors="coerce")
        ok = raw.dropna(subset=["order_date", "unit_price", "quantity", "discount_pct"])
        ok = ok[(ok["product"].notna()) & (ok["unit_price"] >= 0) & (ok["quantity"] > 0)
                & ok["discount_pct"].between(0, 100) & (ok["status"] == "Completed")]
        week = ok[(ok["order_date"] >= "2026-09-21") & (ok["order_date"] <= "2026-09-27")]
        # each order line is rounded to the cent, like an invoice
        expected = round((week["quantity"] * week["unit_price"]
                          * (1 - week["discount_pct"] / 100)).round(2).sum(), 2)
        assert abs(result["current"]["revenue"] - expected) < 0.01

        # the workbook has every sheet, charts, and the invalid rows are listed
        wb = openpyxl.load_workbook(result["xlsx"])
        assert wb.sheetnames == ["Summary", "Daily", "Top Products", "Categories",
                                 "Regions", "Sales Reps", "Weekly Trend", "Skipped Rows"]
        assert len(wb["Summary"]._charts) == 2
        assert wb["Skipped Rows"].max_row == 1 + len(result["skipped"])
        assert wb["Daily"].max_row == 8                 # header + 7 days

        text = result["txt"].read_text()
        assert "Revenue for the week was" in text


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items())
             if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"PASS  {name}")
        except AssertionError as exc:
            failed += 1
            print(f"FAIL  {name}  {exc}")
    print(f"\n{len(tests) - failed}/{len(tests)} tests passed")
    raise SystemExit(1 if failed else 0)
